import os
import shutil
import tempfile
import unittest
import zipfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from latex_to_docx import (
    _apply_direct_formatting,
    _attach_punctuation_to_links,
    _bbl_entry_runs,
    _replace_natbib_citations,
    _create_reference_docx,
    _display_math_blocks,
    _place_references_like_source,
    _restore_title_block,
    _standalone_float_source,
    convert_latex_to_docx,
)


SOURCE = r"""\documentclass{article}
\usepackage{authblk}
\usepackage{caption}
\title{Token sequences}
\author[1]{Ada Lovelace}
\author[2]{Alan Turing}
\author[1]{Grace Hopper}
\affil[1]{Department of Computing Science}
\affil[2]{Department of Mathematics}
\begin{document}
\maketitle
\newcommand{\cell}[1]{\textbf{#1}}
Text.
\bibliography{refs}
\appendix
\section{Extra}
\[
a = b \\[2pt] c
\]
\begin{align*}
x &= y\\[2pt]
z &= w
\end{align*}
\end{document}
"""


def _reference_document(directory):
    path = os.path.join(directory, "reference.docx")
    _create_reference_docx(SOURCE, path)
    return path


def _add_link(paragraph, anchor, text):
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


class LatexLayoutTests(unittest.TestCase):
    def test_reference_docx_uses_black_headings_and_links(self):
        with tempfile.TemporaryDirectory() as directory:
            path = _reference_document(directory)
            with zipfile.ZipFile(path) as archive:
                styles = archive.read("word/styles.xml").decode("utf-8")
            self.assertNotIn('w:themeColor="accent', styles)
            document = Document(path)
            self.assertFalse(document.styles["Title"].font.bold)
            self.assertFalse(document.styles["Hyperlink"].font.underline)

    def test_title_block_joins_authors_and_adds_affiliations(self):
        with tempfile.TemporaryDirectory() as directory:
            document = Document(_reference_document(directory))
            body = document.element.body
            for element in list(body):
                if element.tag != qn("w:sectPr"):
                    body.remove(element)
            for name in ("Ada Lovelace", "Alan Turing", "Grace Hopper"):
                document.add_paragraph(name, style="Author")
            path = os.path.join(directory, "out.docx")
            document.save(path)

            _restore_title_block(SOURCE, path)

            paragraphs = [p for p in Document(path).paragraphs if p.text]
            self.assertEqual(
                paragraphs[0].text, "Ada Lovelace1, Alan Turing2, and Grace Hopper1")
            self.assertTrue(paragraphs[0].runs[1].font.superscript)
            self.assertEqual(paragraphs[1].text, "1Department of Computing Science")
            self.assertEqual(paragraphs[2].text, "2Department of Mathematics")

    def test_punctuation_after_link_moves_into_the_link(self):
        with tempfile.TemporaryDirectory() as directory:
            document = Document()
            paragraph = document.add_paragraph("See Figure ")
            _add_link(paragraph, "fig:a", "1b")
            paragraph.add_run(". Next sentence.")
            path = os.path.join(directory, "out.docx")
            document.save(path)

            _attach_punctuation_to_links(path)

            paragraph = Document(path).paragraphs[0]
            link = paragraph._p.find(qn("w:hyperlink"))
            self.assertEqual("".join(t.text for t in link.iter(qn("w:t"))), "1b.")
            self.assertEqual(paragraph.text, "See Figure 1b. Next sentence.")

    def test_references_move_before_appendix_when_source_puts_them_there(self):
        with tempfile.TemporaryDirectory() as directory:
            document = Document(_reference_document(directory))
            if "Bibliography" not in [style.name for style in document.styles]:
                document.styles.add_style("Bibliography", WD_STYLE_TYPE.PARAGRAPH)
            document.add_paragraph("Funding", style="Heading 1")
            document.add_paragraph("Appendix A Extra", style="Heading 1")
            document.add_paragraph("Lovelace, Ada. 1843. Notes.", style="Bibliography")
            path = os.path.join(directory, "out.docx")
            document.save(path)

            _place_references_like_source(SOURCE, path)

            texts = [p.text for p in Document(path).paragraphs if p.text]
            self.assertEqual(texts[-4:], [
                "Funding", "References", "Lovelace, Ada. 1843. Notes.", "Appendix A Extra"])

    def test_standalone_float_keeps_panels_and_body_macros_but_not_caption(self):
        block = (r"\begin{table}[htbp]" "\n" r"\caption{Main caption}" "\n"
                 r"\begin{tabular}{c}\cell{x}\end{tabular}" "\n" r"\end{table}")
        source = _standalone_float_source(SOURCE, block, "table")
        self.assertIn(r"\newcommand{\cell}[1]{\textbf{#1}}", source)
        self.assertIn(r"\captionsetup{type=table}", source)
        self.assertIn(r"\cell{x}", source)
        self.assertNotIn("Main caption", source)

    def test_display_math_ignores_bracketed_line_spacing(self):
        blocks = _display_math_blocks(SOURCE)
        self.assertEqual(len(blocks), 2)
        self.assertTrue(blocks[0].startswith(r"\["))
        self.assertTrue(blocks[1].startswith(r"\begin{align*}"))

    def test_natbib_citations_are_written_like_the_pdf(self):
        entries = {
            "wang2023": {"short": "Wang et~al.", "year": "2023", "body": ""},
            "wang2024": {"short": "Wang et~al.", "year": "2024", "body": ""},
            "otter2020": {"short": "Otter et~al.", "year": "2020", "body": ""},
        }
        natbib = r"\usepackage[round,authoryear]{natbib}"
        text = r"See \citep{wang2023,wang2024,otter2020} and \citet{otter2020}."
        self.assertEqual(
            _replace_natbib_citations(text, entries, natbib),
            "See (Wang et~al., 2023, 2024; Otter et~al., 2020) and Otter et~al. (2020).")

    def test_bbl_entry_keeps_italics_accents_and_dashes(self):
        body = (r"Hernan Bocaccio and Marisol Dom{\'\i}nguez." "\n"
                r"\newblock Dialects of cardinals." "\n"
                r"\newblock \emph{Ecological Informatics}, 12\penalty0 (3):\penalty0 1--9, 2023.")
        self.assertEqual(_bbl_entry_runs(body), [
            ("Hernan Bocaccio and Marisol Domínguez. Dialects of cardinals. ", False),
            ("Ecological Informatics", True),
            (", 12(3):1\u20139, 2023.", False),
        ])

    def test_style_formatting_is_copied_onto_runs_for_pasting(self):
        with tempfile.TemporaryDirectory() as directory:
            document = Document(_reference_document(directory))
            document.add_paragraph("Results", style="Heading 1")
            document.add_paragraph("Body text.", style="Body Text")
            path = os.path.join(directory, "out.docx")
            document.save(path)

            _apply_direct_formatting(path)

            paragraphs = {p.text: p for p in Document(path).paragraphs}
            heading = paragraphs["Results"].runs[0].font
            body = paragraphs["Body text."]
            self.assertEqual((heading.size.pt, heading.bold, heading.name),
                             (14, True, "Times New Roman"))
            self.assertEqual(body.runs[0].font.size.pt, 12)
            self.assertEqual(body.paragraph_format.line_spacing, 1.5)

    @unittest.skipUnless(shutil.which("pdflatex"), "pdflatex is not installed")
    def test_tables_and_equations_can_be_inserted_as_images(self):
        source = r"""\documentclass{article}
\usepackage{booktabs}
\begin{document}
Text before.
\begin{table}[h]
\caption{Scores}
\begin{tabular}{cc}\toprule a & b\\ \bottomrule\end{tabular}
\end{table}
\[ x^2 + y^2 = z^2 \]
\end{document}
"""
        with tempfile.TemporaryDirectory() as directory:
            source_path = os.path.join(directory, "paper.tex")
            with open(source_path, "w", encoding="utf-8") as source_file:
                source_file.write(source)
            editable = os.path.join(directory, "editable.docx")
            pictures = os.path.join(directory, "pictures.docx")
            convert_latex_to_docx(source_path, editable)
            convert_latex_to_docx(source_path, pictures, tables_as_images=True,
                                  equations_as_images=True)

            self.assertGreaterEqual(len(Document(editable).tables), 1)
            self.assertEqual(len(Document(editable).inline_shapes), 0)
            converted = Document(pictures)
            self.assertEqual(len(converted.inline_shapes), 2)
            caption = next(p for p in converted.paragraphs if p.text.startswith("Table 1:"))
            picture = caption._p.getnext()
            self.assertIsNotNone(next(picture.iter(qn("w:drawing")), None))
            self.assertEqual(len(converted.element.body.findall(".//" + qn("m:oMathPara"))), 0)


if __name__ == "__main__":
    unittest.main()
